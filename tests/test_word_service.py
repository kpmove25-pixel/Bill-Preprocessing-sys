"""Word 文档生成服务测试。"""
import os

from PIL import Image

from app.services.word_service import generate_buyer_word


def _make_jpg(path):
    Image.new("RGB", (100, 50), (255, 255, 255)).save(path, "JPEG")


def test_generate_buyer_word(tmp_path):
    jpg1 = str(tmp_path / "a.jpg")
    jpg2 = str(tmp_path / "b.jpg")
    _make_jpg(jpg1)
    _make_jpg(jpg2)

    invoices = [
        {
            "invoice_date": "2026-08-06", "invoice_code": "", "invoice_no": "26412000002974770136",
            "seller_name": "河南正开元机械科技有限公司", "amount": "46800.00",
            "amount_excluding_tax": "41415.92", "archived_jpg": jpg1, "jpg_path": jpg1,
        },
        {
            "invoice_date": "2026-08-05", "invoice_code": "", "invoice_no": "26432000001835823166",
            "seller_name": "湘潭华进重装科技股份有限公司", "amount": "200000.00",
            "amount_excluding_tax": "176991.15", "archived_jpg": jpg2, "jpg_path": jpg2,
        },
    ]

    out = str(tmp_path / "测试公司.docx")
    generate_buyer_word("测试公司", "911", invoices, out)
    assert os.path.exists(out)

    from docx import Document

    doc = Document(out)
    table = doc.tables[0]
    headers = [c.text for c in table.rows[0].cells]
    assert headers == ["序号", "开票时间", "发票代码", "发票号码", "收款方", "金额", "金额（不含税）"]
    assert len(table.rows) == 3  # 表头 + 2 数据行
    row1 = [c.text for c in table.rows[1].cells]
    assert row1[3] == "26412000002974770136"
    assert len(doc.inline_shapes) == 2  # 2 张图片


def test_generate_buyer_word_skips_missing_image(tmp_path):
    invoices = [{
        "invoice_date": "2026-08-06", "invoice_code": "", "invoice_no": "X",
        "seller_name": "S", "amount": "1", "amount_excluding_tax": "1",
        "archived_jpg": str(tmp_path / "不存在.jpg"), "jpg_path": "",
    }]
    out = str(tmp_path / "t.docx")
    generate_buyer_word("T", "", invoices, out)

    from docx import Document

    doc = Document(out)
    assert len(doc.inline_shapes) == 0  # 图片不存在则跳过
