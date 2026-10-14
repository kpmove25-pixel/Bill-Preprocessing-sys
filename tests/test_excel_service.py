"""Excel 生成测试。"""
import os

from openpyxl import load_workbook

from app.services import excel_service


def test_generate_buyer_excel(tmp_path):
    invs = [{
        "invoice_code": "011001900111",
        "invoice_no": "12345678",
        "invoice_type": "增值税专用发票",
        "buyer_name": "测试公司",
        "buyer_tax_no": "91310000XXXX",
        "amount": "1130.00",
        "amount_excluding_tax": "1000.00",
        "tax_amount": "130.00",
        "tax_rate": "13.00%",
        "invoice_date": "2023-05-15",
        "seller_name": "销售公司",
        "seller_tax_no": "",
        "check_code": "",
        "drawer": "",
        "original_filename": "a.pdf",
        "jpg_path": "/tmp/a.jpg",
        "status": "confirmed",
        "remark": "",
    }]
    path = str(tmp_path / "测试公司_发票汇总.xlsx")
    excel_service.generate_buyer_excel(invs, path)
    assert os.path.exists(path)
    wb = load_workbook(path)
    ws = wb.active
    # 表头 + 1 行数据
    assert ws.max_row == 2
    assert ws.cell(row=1, column=1).value == "序号"
    assert ws.cell(row=2, column=3).value == "011001900111"
    assert ws.cell(row=2, column=4).value == "12345678"


def test_build_buyer_excel_name():
    name = excel_service.build_buyer_excel_name("测试公司", "91310000XXXX")
    assert name == "测试公司_91310000XXXX_发票汇总.xlsx"
    name2 = excel_service.build_buyer_excel_name("测试公司", "")
    assert name2 == "测试公司_发票汇总.xlsx"
